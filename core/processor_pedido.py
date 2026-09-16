import os
import re
from typing import Optional, Tuple, Dict, Any
import openpyxl
from .models import ResultadoPedido
from .file_utils import mover_arquivo_seguro
from .excel_utils import limpar_coluna_observacao, garantir_planilha_base_existe

def normalizar_nota(s: str) -> str:
    """
    Normaliza o numero da nota/remessa para comparacao.
    Remove extensao, prefixos e sufixos apos o hifen e zeros a esquerda.
    """
    s = os.path.splitext(str(s))[0].strip()
    antes_hifen = s.split("-")[0]
    nums_antes = re.findall(r'\d+', antes_hifen)
    if nums_antes:
        return nums_antes[-1].lstrip("0") or "0"
    numeros = re.findall(r'\d+', s)
    if numeros:
        return numeros[0].lstrip("0") or "0"
    return s.lstrip("0") or "0"


def ler_planilha_pedido(caminho: Optional[str] = None):
    """Le a aba 'PEDIDO' da planilha."""
    if not caminho or not os.path.exists(caminho):
        caminho = garantir_planilha_base_existe()
    
    wb = openpyxl.load_workbook(caminho)
    aba = wb["PEDIDO"] if "PEDIDO" in wb.sheetnames else wb.worksheets[1]
    
    mapa = {}
    for row in aba.iter_rows(min_row=2):
        val_a, val_b = row[0].value, row[1].value
        if val_a is not None:
            chave = normalizar_nota(str(val_a).strip())
            mapa[chave] = (str(val_b).strip() if val_b is not None else "", row[0].row)
    return mapa, wb, aba


def analisar_pedido(pasta: str, caminho_excel: Optional[str] = None) -> ResultadoPedido:
    """Analisa os arquivos de remessa cruzando com a aba de pedidos."""
    mapa, _, _ = ler_planilha_pedido(caminho_excel)
    
    if not os.path.exists(pasta):
        return ResultadoPedido()
        
    arquivos = [f for f in os.listdir(pasta) if f.lower().endswith(".pdf")]
    resultado = ResultadoPedido()
    encontrados = set()

    for nome in arquivos:
        chave = normalizar_nota(nome)
        encontrados.add(chave)
        if chave not in mapa:
            resultado.nao_listado.append((chave, nome))
            continue
        pedido, _ = mapa[chave]
        pedido_up = pedido.upper()
        if "PROPAG" in pedido_up:
            resultado.propag.append((chave, nome))
        elif "ETI" in pedido_up:
            resultado.eti.append((chave, nome))
        else:
            resultado.sem_destino.append((chave, nome))

    for chave, (_, linha) in mapa.items():
        if chave not in encontrados:
            resultado.nao_achado.append((chave, linha))

    return resultado


def mover_pedido(pasta: str, pasta_eti: str, pasta_propag: str, caminho_excel: Optional[str],
                 resultado: ResultadoPedido, callback_log, callback_status):
    """Move os arquivos de pedidos para ETI ou PROPAG e atualiza a planilha."""
    erros = 0
    movidos_eti = set()
    movidos_propag = set()
    falhas_mover = {}

    def _mover(lista, destino, categoria):
        nonlocal erros
        count = 0
        for chave, nome in lista:
            try:
                _, aviso = mover_arquivo_seguro(os.path.join(pasta, nome), destino)
                if aviso:
                    callback_log(aviso)
                if categoria == "eti":
                    movidos_eti.add(chave)
                else:
                    movidos_propag.add(chave)
                count += 1
            except Exception as e:
                callback_log(f"Erro ao mover {nome}: {e}")
                falhas_mover[chave] = str(e)
                erros += 1
        return count

    me = _mover(resultado.eti,    pasta_eti,    "eti")
    mp = _mover(resultado.propag, pasta_propag, "propag")

    try:
        _, wb, aba = ler_planilha_pedido(caminho_excel)
        ultima_dados = limpar_coluna_observacao(aba, col_obs=3)

        sem_destino_chaves = {chave for chave, _ in resultado.sem_destino}
        nao_achados_chaves = {chave for chave, _ in resultado.nao_achado}

        for r in range(2, ultima_dados + 1):
            cell_a = aba.cell(row=r, column=1).value
            cell_b = aba.cell(row=r, column=2).value
            if cell_a is None:
                continue
            chave = normalizar_nota(str(cell_a))
            pedido_texto = str(cell_b).strip() if cell_b is not None else ""

            if chave in movidos_eti:
                aba.cell(row=r, column=3, value="OK - Movido para ETI")
            elif chave in movidos_propag:
                aba.cell(row=r, column=3, value="OK - Movido para PROPAG")
            elif chave in falhas_mover:
                aba.cell(row=r, column=3, value=f"Erro ao mover ({falhas_mover[chave]})")
            elif chave in sem_destino_chaves:
                aba.cell(row=r, column=3, value=f"Sem destino definido (Pedido: {pedido_texto})")
            elif chave in nao_achados_chaves:
                aba.cell(row=r, column=3, value="Nao encontrado na pasta de remessas")
            else:
                aba.cell(row=r, column=3, value="Nao encontrado na pasta de remessas")

        for i, (chave, nome_arq) in enumerate(resultado.nao_listado, start=1):
            aba.cell(row=ultima_dados + i, column=3,
                     value=f"Nao listado na planilha: {nome_arq} (remessa {chave})")

        wb.save(caminho_excel or garantir_planilha_base_existe())
    except PermissionError:
        callback_log("ERRO: Nao foi possivel salvar a planilha. Feche o arquivo Excel e tente novamente.")
        erros += 1
    except Exception as e:
        callback_log(f"ERRO ao salvar planilha: {e}")
        erros += 1

    callback_status(me, mp, erros)