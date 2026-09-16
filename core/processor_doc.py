import os
import re
from typing import Optional, Set, Tuple
import openpyxl
from .models import ResultadoAnalise
from .file_utils import mover_arquivo_seguro
from .excel_utils import limpar_coluna_observacao, garantir_planilha_base_existe

def extrair_codigo(nome_arquivo: str, codigos_planilha: Optional[Set[str]] = None) -> str:
    """
    Extrai o codigo da nota fiscal a partir do nome do arquivo.
    Suporta correspondencia direta se codigos_planilha forem fornecidos.
    """
    nome_sem_ext = os.path.splitext(nome_arquivo)[0].strip()
    if codigos_planilha:
        for cod in sorted(codigos_planilha, key=len, reverse=True):
            if cod in nome_sem_ext:
                return cod
    match_hifen = re.findall(r'\b\d+(?:-\d+)?\b', nome_sem_ext)
    if match_hifen:
        return match_hifen[-1]
    partes = re.split(r'[\s_]+', nome_sem_ext)
    if partes:
        return partes[-1].strip()
    return nome_sem_ext


def ler_planilha_doc(caminho: Optional[str] = None):
    """Le a aba 'XML' da planilha de notas."""
    if not caminho or not os.path.exists(caminho):
        caminho = garantir_planilha_base_existe()
    
    wb = openpyxl.load_workbook(caminho)
    aba = wb["XML"] if "XML" in wb.sheetnames else wb.active
    
    fatura, remessa = {}, {}
    for row in aba.iter_rows(min_row=2):
        val_a, val_b, n = row[0].value, row[1].value, row[0].row
        if val_a is not None:
            fatura[str(val_a).strip()] = n
        if val_b is not None:
            remessa[str(val_b).strip()] = n
    return fatura, remessa, wb, aba


def analisar_documentos(pasta: str, caminho_excel: Optional[str] = None, extensao: str = ".xml") -> ResultadoAnalise:
    """Analisa os arquivos de uma pasta cruzando com as colunas de fatura e remessa."""
    fatura_map, remessa_map, _, _ = ler_planilha_doc(caminho_excel)
    todos_codigos = set(fatura_map.keys()) | set(remessa_map.keys())
    
    if not os.path.exists(pasta):
        return ResultadoAnalise()
        
    arquivos = [f for f in os.listdir(pasta) if f.lower().endswith(extensao)]
    resultado = ResultadoAnalise()
    encontrados = set()

    for nome in arquivos:
        codigo = extrair_codigo(nome, todos_codigos)
        encontrados.add(codigo)
        em_a, em_b = codigo in fatura_map, codigo in remessa_map
        if em_a and em_b:
            resultado.conflito.append((codigo, fatura_map[codigo], remessa_map[codigo]))
        elif em_a:
            resultado.fatura.append((codigo, nome))
        elif em_b:
            resultado.remessa.append((codigo, nome))
        else:
            resultado.nao_listado.append((codigo, nome))

    for codigo, linha in fatura_map.items():
        if codigo not in encontrados:
            resultado.nao_achado.append((codigo, "A", linha))
    for codigo, linha in remessa_map.items():
        if codigo not in encontrados:
            resultado.nao_achado.append((codigo, "B", linha))

    return resultado


def mover_documentos(pasta: str, pasta_fatura: str, pasta_remessa: str, caminho_excel: Optional[str],
                     resultado: ResultadoAnalise, callback_log, callback_status, extensao: str = ".xml"):
    """Move os arquivos de documentos e atualiza as observacoes na planilha."""
    erros = 0
    movidos_fatura = set()
    movidos_remessa = set()
    falhas_mover = {}

    def _mover_lote(lista, destino, categoria):
        nonlocal erros
        count = 0
        for codigo, nome in lista:
            try:
                _, aviso = mover_arquivo_seguro(os.path.join(pasta, nome), destino)
                if aviso:
                    callback_log(aviso)
                if categoria == "fatura":
                    movidos_fatura.add(codigo)
                else:
                    movidos_remessa.add(codigo)
                count += 1
            except Exception as e:
                callback_log(f"Erro ao mover {nome}: {e}")
                falhas_mover[codigo] = str(e)
                erros += 1
        return count

    mf = _mover_lote(resultado.fatura,  pasta_fatura,  "fatura")
    mr = _mover_lote(resultado.remessa, pasta_remessa, "remessa")

    try:
        _, _, wb, aba = ler_planilha_doc(caminho_excel)
        ext = extensao.upper().lstrip(".")
        ultima_dados = limpar_coluna_observacao(aba, col_obs=3)

        conflitos_set = {c[0] for c in resultado.conflito}
        nao_achados_map = {(codigo, col): linha for codigo, col, linha in resultado.nao_achado}

        for r in range(2, ultima_dados + 1):
            cell_a = aba.cell(row=r, column=1).value
            cell_b = aba.cell(row=r, column=2).value
            cod_a = str(cell_a).strip() if cell_a is not None else None
            cod_b = str(cell_b).strip() if cell_b is not None else None

            partes_obs = []

            if cod_a:
                if cod_a in conflitos_set:
                    partes_obs.append("FATURA: Conflito (presente em FATURA e REMESSA)")
                elif cod_a in movidos_fatura:
                    partes_obs.append("FATURA: OK")
                elif cod_a in falhas_mover:
                    partes_obs.append(f"FATURA: Erro ao mover ({falhas_mover[cod_a]})")
                elif (cod_a, "A") in nao_achados_map:
                    partes_obs.append(f"FATURA: Nao encontrado na pasta {ext}")
                else:
                    partes_obs.append("FATURA: Nao encontrado")

            if cod_b:
                if cod_b in conflitos_set:
                    partes_obs.append("REMESSA: Conflito (presente em FATURA e REMESSA)")
                elif cod_b in movidos_remessa:
                    partes_obs.append("REMESSA: OK")
                elif cod_b in falhas_mover:
                    partes_obs.append(f"REMESSA: Erro ao mover ({falhas_mover[cod_b]})")
                elif (cod_b, "B") in nao_achados_map:
                    partes_obs.append(f"REMESSA: Nao encontrado na pasta {ext}")
                else:
                    partes_obs.append("REMESSA: Nao encontrado")

            if partes_obs:
                aba.cell(row=r, column=3, value=" | ".join(partes_obs))

        for i, (codigo, nome_arq) in enumerate(resultado.nao_listado, start=1):
            aba.cell(row=ultima_dados + i, column=3,
                     value=f"Nao listado na planilha: {nome_arq} (codigo {codigo})")

        wb.save(caminho_excel or garantir_planilha_base_existe())
    except PermissionError:
        callback_log("ERRO: Nao foi possivel salvar a planilha. Feche o arquivo Excel e tente novamente.")
        erros += 1
    except Exception as e:
        callback_log(f"ERRO ao salvar planilha: {e}")
        erros += 1

    callback_status(mf, mr, erros)